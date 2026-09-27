from pathlib import Path
import argparse
import csv
import json
import math
import hashlib
import openpyxl

parser = argparse.ArgumentParser(description="Summarize a separately obtained measured-MTF workbook.")
parser.add_argument("workbook", type=Path, help="Path to a license-compatible source workbook.")
parser.add_argument("--out-dir", type=Path, default=Path("."), help="Directory for JSON/CSV outputs.")
args = parser.parse_args()
DATA = args.workbook
OUT = args.out_dir
OUT.mkdir(parents=True, exist_ok=True)


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def pairs(sheet):
    result = []
    for row in range(2, sheet.max_row + 1):
        power = sheet.cell(row, 5).value
        mtf = sheet.cell(row, 6).value
        if isinstance(power, (int, float)) and isinstance(mtf, (int, float)):
            if math.isfinite(float(power)) and math.isfinite(float(mtf)):
                result.append((float(power), float(mtf)))
    return result


wb = openpyxl.load_workbook(DATA, data_only=True, read_only=True)
records = []
for sheet in wb.worksheets:
    data = pairs(sheet)
    if not data:
        records.append({"sheet": sheet.title, "status": "fail", "error": "no finite Power/MTF pairs"})
        continue
    peak_power, peak_mtf = max(data, key=lambda item: item[1])
    row = {
        "sheet": sheet.title,
        "status": "pass",
        "n_pairs": len(data),
        "power_min_D": min(p for p, _ in data),
        "power_max_D": max(p for p, _ in data),
        "mtf_min": min(m for _, m in data),
        "mtf_max": max(m for _, m in data),
        "peak_power_D": peak_power,
        "peak_mtf": peak_mtf,
    }
    for threshold in (0.10, 0.20):
        eligible = [p for p, m in data if m >= threshold]
        row[f"power_width_at_{int(threshold*100)}pct_D"] = (max(eligible) - min(eligible)) if eligible else 0.0
        row[f"n_at_{int(threshold*100)}pct"] = len(eligible)
    records.append(row)

summary = {
    "status": "completed",
    "dataset": "Son et al. 2020 PLOS ONE supplementary raw pixel-intensity and MTF data",
    "measurement_context": "four IOL models, 3.0 mm and 4.5 mm pupils, OptiSpheric IOL PRO II, laboratory measurements",
    "source_file": DATA.name,
    "source_sha256": sha(DATA),
    "records": records,
    "limitation": "Measured optical data but not an imaging-lens prescription; used as an external MTF measurement benchmark, not as a direct candidate-ranking replication.",
}
(OUT / "public_measured_mtf_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
with (OUT / "public_measured_mtf_summary.csv").open("w", newline="", encoding="utf-8") as handle:
    fields = sorted({key for row in records for key in row})
    writer = csv.DictWriter(handle, fieldnames=fields)
    writer.writeheader()
    writer.writerows(records)
print(json.dumps({"status": "completed", "records": records}, indent=2))
